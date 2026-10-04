import json
from pathlib import Path

from docx import Document
from docx.shared import Inches


def create_document(progress_file, output_file):
    progress_path = Path(progress_file)

    if not progress_path.exists():
        raise RuntimeError("No progress file found.")

    data = json.loads(progress_path.read_text(encoding="utf-8"))

    document = Document()

    document.add_heading("W3Schools Exercises", level=0)

    document.add_paragraph(
        "Automatically completed W3Schools HTML, CSS and JavaScript exercises."
    )

    entries = data.get("completed", [])

    current_subject = None

    for number, entry in enumerate(entries, start=1):
        subject = entry["subject"]
        category = entry["category"]
        question = entry["question"]
        screenshot = entry["screenshot"]

        if subject != current_subject:
            document.add_page_break()

            document.add_heading(subject, level=1)

            current_subject = subject

        document.add_heading(f"{category} - Question {question}", level=2)

        document.add_paragraph(f"Exercise #{number}")

        screenshot_path = Path(screenshot)

        if screenshot_path.exists():
            document.add_picture(str(screenshot_path), width=Inches(6.5))
        else:
            document.add_paragraph("Screenshot missing.")

        if number != len(entries):
            document.add_page_break()

    document.save(output_file)
